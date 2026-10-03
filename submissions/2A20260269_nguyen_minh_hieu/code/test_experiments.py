"""Kiểm tra tự viết cho experiments.py bằng dữ liệu giả (không cần GPU, không cần dataset):
    python -m unittest test_experiments -v
"""
import json
import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
import torch
from openpyxl import load_workbook

import experiments as E
import train
from train import Config, ev

K = 9


def fake_summary(cfg, f1, per_class=None):
    d = train.run_dir(cfg)
    d.mkdir(parents=True, exist_ok=True)
    s = {"exp_id": cfg.exp_id, "seed": cfg.seed, "backbone": cfg.backbone, "weights_tag": f"{cfg.backbone}.tag",
         "params_m": 23.5, "gmacs": 4.1, "best_epoch": 9, "epochs": cfg.epochs, "val_macro_f1": f1,
         "val_top1": f1 + 0.02, "val_ece": 0.03, "val_f1_per_class": per_class or [f1] * K,
         "train_time_per_epoch_s": 40.0}
    (d / "summary.json").write_text(json.dumps(s), encoding="utf-8")


class TestSpecs(unittest.TestCase):
    def test_each_training_spec_changes_only_listed_fields(self):
        base = Config(backbone="resnet50")
        for cfg, (eid, axis, desc, over) in zip(E.training_configs("resnet50", {}), E.TRAINING):
            changed = {k for k, v in vars(cfg).items() if v != getattr(base, k)} - {"exp_id", "desc"}
            self.assertEqual(changed, set(over), eid)
            self.assertIn(axis, "ABCDEFG")
        ids = [s[0] for s in E.TRAINING]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertGreaterEqual(len({s[1] for s in E.TRAINING}), 3)

    def test_fixed_input_backbones_skip_resolution_change(self):
        ids = [c.exp_id for c in E.training_configs("swin_tiny", {})]
        self.assertNotIn("T15", ids)
        self.assertIn("T15", [c.exp_id for c in E.training_configs("resnet50", {})])

    def test_backbones_cover_required_families(self):
        keys = [k for _, k in E.BACKBONES]
        self.assertGreaterEqual(len(keys), 5)
        for needed in ("resnet50", "convnext_tiny", "swin_tiny", "mobilenetv3"):
            self.assertIn(needed, keys)


class TestSelection(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.common = {"out_dir": self.tmp.name}
        self.t00 = [E.t00_config("resnet50", s, self.common) for s in (0, 1, 2)]
        for cfg, f1 in zip(self.t00, (0.900, 0.904, 0.896)):          # std = 0.004
            fake_summary(cfg, f1)
        self.t = E.training_configs("resnet50", self.common)
        f1 = {"T01": 0.60, "T02": 0.70, "T03": 0.902, "T04": 0.895, "T05": 0.930, "T06": 0.915, "T07": 0.90,
              "T08": 0.910, "T09": 0.906, "T10": 0.89, "T11": 0.899, "T12": 0.88, "T13": 0.903, "T14": 0.901,
              "T15": 0.92, "T16": 0.925}
        for cfg in self.t:
            fake_summary(cfg, f1[cfg.exp_id])

    def tearDown(self):
        self.tmp.cleanup()

    def test_noise_and_table(self):
        self.assertAlmostEqual(E.noise_std(self.t00), 0.004, places=6)
        table = E.training_table(self.t00, self.t)
        self.assertEqual(len(table), 3 + 16)
        row = table.set_index("exp_id").loc["T05"]
        self.assertAlmostEqual(row["delta_vs_t00"], 0.030, places=6)
        self.assertIn("tốt hơn", row["note"])
        self.assertIn("không phân biệt", table.set_index("exp_id").loc["T03", "note"])   # +0.002 < std
        self.assertIn("kém hơn", table.set_index("exp_id").loc["T12", "note"])

    def test_combo_takes_best_per_axis_above_noise(self):
        table = E.training_table(self.t00, self.t)
        over, used = E.combo_overrides(table, E.noise_std(self.t00))
        # B: T05 (+0.030), C: T08 (+0.010), G: T16 (+0.025); D, E, F không vượt nhiễu; A không tham gia
        self.assertEqual(sorted(used), ["T05", "T08", "T16"])
        self.assertEqual(over, {"aug": "mildcrop", "loss": "ls", "label_smoothing": 0.1, "epochs": 20})

    def test_combo_falls_back_to_top_two(self):
        for cfg in self.t:
            fake_summary(cfg, 0.899 if cfg.exp_id not in ("T05", "T08") else 0.9005)
        table = E.training_table(self.t00, self.t)
        _, used = E.combo_overrides(table, E.noise_std(self.t00))
        self.assertEqual(sorted(used), ["T05", "T08"])

    def test_pick_final_and_backbone(self):
        combo = Config(exp_id=E.COMBO_ID, backbone="resnet50", **self.common)
        fake_summary(combo, 0.940)
        table = E.training_table(self.t00, self.t, combo, "T05 + T08 + T16")
        self.assertEqual(E.pick_final(table), "T20")
        self.assertEqual(E.overrides_of("T20", {"aug": "mildcrop"}), {"aug": "mildcrop"})
        self.assertEqual(E.overrides_of("T00"), {})
        self.assertEqual(E.overrides_of("T14"), {"ema_decay": 0.998})

        bcfgs = E.backbone_configs(self.common)
        for cfg, f1 in zip(bcfgs, (0.90, 0.91, 0.95, 0.93, 0.95, 0.88, 0.87)):
            fake_summary(cfg, f1)
        bt = E.backbone_table(bcfgs, {"resnet50": 7.5})
        self.assertEqual(len(bt), 7)
        self.assertEqual(E.pick_backbone(bt), "convnext_tiny")
        self.assertEqual(bt.set_index("backbone").loc["resnet50", "latency_b1_ms"], 7.5)


class TestViews(unittest.TestCase):
    def test_views_shapes_224(self):
        views, methods = E.build_views(224, fixed_input=False)
        x = torch.randn(2, 3, 256, 256)
        self.assertEqual(views["c"](x).shape[-1], 224)
        torch.testing.assert_close(views["c"](x), x[..., 16:240, 16:240])
        torch.testing.assert_close(views["c_f"](x), torch.flip(x[..., 16:240, 16:240], [-1]))
        torch.testing.assert_close(views["crop4"](x), views["c"](x))
        self.assertIs(views["r256"](x), x)
        self.assertEqual(views["r320"](x).shape[-1], 320)
        ids = [m["exp_id"] for m in methods]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertGreaterEqual(len(ids) - 1, 4)                     # >= 4 phương pháp ngoài I00
        for m in methods:
            self.assertTrue(set(m["views"]) <= set(views))
            self.assertEqual(sum(n for _, n in m["parts"]), len(m["views"]))

    def test_fixed_input_has_no_resolution_sweep(self):
        views, methods = E.build_views(224, fixed_input=True)
        self.assertFalse(any(m["exp_id"].startswith("I04") for m in methods))
        self.assertTrue(all(v(torch.randn(1, 3, 256, 256)).shape[-1] == 224 for v in views.values()))

    def test_views_256(self):
        views, methods = E.build_views(256, fixed_input=False)
        x = torch.randn(1, 3, 256, 256)
        self.assertIs(views["c"](x), x)
        self.assertEqual(views["crop0"](x).shape[-1], 256)
        self.assertEqual([m["exp_id"] for m in methods if m["exp_id"].startswith("I04")],
                         ["I04a", "I04af", "I04b", "I04bf"])

    def test_log_temperature_matches_standard_ts_for_single_view(self):
        rng = np.random.default_rng(0)
        logits = rng.normal(size=(500, K)) * 4
        y = rng.integers(0, K, 500)
        p = torch.from_numpy(logits).softmax(1).numpy()
        import inference
        self.assertAlmostEqual(E.fit_log_temperature(p, y), inference.fit_temperature(logits, y), places=4)
        np.testing.assert_allclose(E.apply_log_temperature(p, 2.0), inference.apply_temperature(logits, 2.0),
                                   atol=1e-8)


class TestTablesAndXlsx(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        self.labels, self.preds = root / "labels", root / "predictions"
        self.labels.mkdir()
        rng = np.random.default_rng(0)
        n = 300
        for split in ("val", "test"):
            y = rng.integers(0, K, n)
            names = [f"{split}{i}.jpg" for i in range(n)]
            pd.DataFrame({"Filename": names, "Label": y}).to_csv(self.labels / f"{split}_subset0.csv", index=False)
            for tag, sharp in (("F01", 6.0), ("T00", 3.0), ("F01rt", 5.0), ("F01_uncal", 9.0)):
                for seed in range(3):
                    z = rng.normal(size=(n, K))
                    z[np.arange(n), y] += sharp
                    p = torch.from_numpy(z).softmax(1).numpy()
                    ev.save_predictions(self.preds / f"{tag}_seed{seed}_{split}.csv", names, y, p)

    def tearDown(self):
        self.tmp.cleanup()

    def test_final_and_per_class_match_eval(self):
        final = E.final_table(self.preds, self.labels, {"F01": "chung kết", "T00": "mốc", "F01rt": "thời gian thực"})
        self.assertEqual(len(final), 3 * 4)
        g = ev.load_group(str(self.preds / "F01_seed*_test.csv"), str(self.labels / "test_subset0.csv"))
        mean_row = final[(final["exp_id"] == "F01") & final["seed"].astype(str).str.startswith("mean")].iloc[0]
        self.assertAlmostEqual(mean_row["test_macro_f1"], g.summary["macro_f1"][0], places=12)
        self.assertIn("±", mean_row["mean_pm_std"])
        seed1 = final[(final["exp_id"] == "F01") & (final["seed"] == 1)].iloc[0]
        self.assertAlmostEqual(seed1["test_top1"], g.metrics[1]["top1"], places=12)
        # glob của F01 không được lẫn F01_uncal hay F01rt
        self.assertEqual(len(g.preds), 3)

        pc = E.per_class_table(self.preds, self.labels, ["T00", "F01"])
        self.assertEqual(len(pc), 2 * K)
        row = pc[(pc["exp_id"] == "F01") & (pc["class"] == "Snake Weed")].iloc[0]
        self.assertAlmostEqual(row["recall"], g.summary["recall"][0][7], places=12)
        self.assertEqual(E.confusion_sum(self.preds, self.labels, "F01").sum(), 900)

    def test_summary_and_xlsx(self):
        common = {"out_dir": self.tmp.name}
        bcfgs = E.backbone_configs(common)
        for cfg, f1 in zip(bcfgs, (0.90, 0.91, 0.95, 0.93, 0.94, 0.88, 0.87)):
            fake_summary(cfg, f1)
        t00 = [E.t00_config("convnext_tiny", s, common) for s in (0, 1, 2)]
        for cfg, f1 in zip(t00, (0.95, 0.951, 0.949)):
            fake_summary(cfg, f1)
        tcfgs = E.training_configs("convnext_tiny", common)
        for i, cfg in enumerate(tcfgs):
            fake_summary(cfg, 0.90 + 0.004 * i)
        backbones = E.backbone_table(bcfgs, {})
        training = E.training_table(t00, tcfgs)
        inf = pd.DataFrame([
            {"exp_id": "I00", "method": "1 view", "model": "m", "k": 1, "val_macro_f1": 0.96, "val_top1": 0.97,
             "val_ece": 0.02, "p50_ms": 8.0, "p95_ms": 9.0, "p99_ms": 10.0, "images_per_s_b32": 400.0,
             "single_model": True, "note": "", "rel_cost": 1.0},
            {"exp_id": "I01", "method": "lật", "model": "m", "k": 2, "val_macro_f1": 0.965, "val_top1": 0.972,
             "val_ece": 0.02, "p50_ms": 16.0, "p95_ms": 18.0, "p99_ms": 20.0, "images_per_s_b32": 200.0,
             "single_model": True, "note": "", "rel_cost": 2.0},
            {"exp_id": "I05", "method": "ensemble", "model": "m+n", "k": 2, "val_macro_f1": 0.97, "val_top1": 0.975,
             "val_ece": 0.02, "p50_ms": 150.0, "p95_ms": 160.0, "p99_ms": 170.0, "images_per_s_b32": 100.0,
             "single_model": False, "note": "", "rel_cost": 18.75}])
        self.assertEqual(E.pick_method(inf), "I01")                  # ensemble không phải ứng viên một mô hình
        self.assertEqual(E.pick_realtime(inf)["exp_id"], "I00")
        self.assertIsNone(E.pick_realtime(inf.assign(p95_ms=500.0)))

        final = E.final_table(self.preds, self.labels, {"F01": "chung kết", "T00": "mốc"})
        summary = E.summary_table(backbones, training, inf, final)
        self.assertEqual(summary.iloc[0]["step"], "CHUNG KẾT (test)")
        self.assertIn("T00", summary["exp_id"].values)
        self.assertLessEqual(len(summary), 2 + 10 + 1)

        sheets = {"Backbones": backbones, "Training": training, "Inference": inf, "Final": final,
                  "PerClass": E.per_class_table(self.preds, self.labels, ["T00", "F01"]),
                  "Latency": pd.DataFrame([{"config": "I00", "gpu": "T4", "dtype": "fp32", "batch": 1,
                                            "fused_bn": False, "p50": 8.0, "p95": 9.0, "p99": 10.0,
                                            "images_per_s": 125.0}]),
                  "Summary": summary}
        path = E.write_xlsx(Path(self.tmp.name) / "results.xlsx", sheets)
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, list(sheets))
        ws = wb["Backbones"]
        self.assertEqual(ws.freeze_panes, "A2")
        header = [c.value for c in ws[1]]
        for col in ("exp_id", "backbone", "tag trọng số (timm)", "#tham số (M)", "GMAC", "macro-F1 val",
                    "top-1 val", "thời gian train/epoch (s)", "độ trễ batch-1 p50 (ms)", "ghi chú"):
            self.assertIn(col, header)
        best_row = 2 + int(backbones["val_macro_f1"].astype(float).values.argmax())
        self.assertEqual(ws.cell(best_row, 1).fill.start_color.rgb[-6:], "FFF2CC")
        self.assertEqual(ws.cell(2, header.index("macro-F1 val") + 1).number_format, "0.0000")
        for name in sheets:                                          # mọi cột đã có tên tiếng Việt/đơn vị
            unmapped = [c for c in sheets[name].columns if c not in E.COLUMNS]
            self.assertEqual(unmapped, [], name)

    def test_config_description(self):
        cfg = Config(exp_id="F01", backbone="convnext_tiny", aug="mildcrop", epochs=20, seed=2)
        self.assertEqual(E.config_description(cfg, "TTA lật"), "convnext_tiny | aug=mildcrop, epochs=20 | TTA lật")
        self.assertEqual(E.config_description(Config(), "1 view"), "resnet50 | công thức nền | 1 view")


if __name__ == "__main__":
    unittest.main()
