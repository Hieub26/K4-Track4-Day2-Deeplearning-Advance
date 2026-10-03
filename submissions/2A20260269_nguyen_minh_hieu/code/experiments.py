"""experiments.py - danh sách thí nghiệm, chọn cấu hình trên VAL, nghiên cứu suy luận, bảng và results.xlsx.

Notebook lab_day2.ipynb chỉ gọi các hàm ở đây; mọi lần huấn luyện vẫn đi qua train.run(Config).
Mọi lựa chọn tự động (backbone, tổ hợp công thức, phương pháp suy luận) chỉ dùng số liệu VAL.
Test chỉ được đọc trong final_predictions(), đúng một lần cho mỗi seed (file đã có thì không chạy lại).

Quy ước exp_id: B01.. backbone · T00 nền, T01.. công thức, T20 tổ hợp · I00.. suy luận · F01 chung kết,
F01_uncal = F01 chưa temperature scaling, F01rt = cấu hình thời gian thực (cùng trọng số F01, 1 view).
"""
from __future__ import annotations

import copy
import dataclasses
import json
from pathlib import Path

import numpy as np
import pandas as pd
import torch
import torch.nn.functional as F

import benchmark
import dataset
import inference
import model as model_lib
import train
from train import Config, ev

# (exp_id, khoá backbone trong model.SUGGESTED_BACKBONES)
BACKBONES = [
    ("B01", "resnet50"), ("B02", "resnext50"), ("B03", "convnext_tiny"), ("B04", "deit_small"),
    ("B05", "swin_tiny"), ("B06", "efficientnet_b0"), ("B07", "mobilenetv3"),
]
FIXED_INPUT = {"deit_small", "swin_tiny"}   # timm cố định đầu vào 224: không đổi độ phân giải được

# (exp_id, trục GUIDE mục 3, mô tả ngắn, các field khác T00). Mỗi dòng chỉ khác T00 MỘT yếu tố.
TRAINING = [
    ("T01", "A", "scratch", {"init": "scratch"}),
    ("T02", "A", "frozen", {"init": "frozen"}),
    ("T03", "B", "color", {"aug": "color"}),
    ("T04", "B", "trivial", {"aug": "trivial"}),
    ("T05", "B", "mildcrop", {"aug": "mildcrop"}),
    ("T06", "B", "cutmix", {"mix": "cutmix"}),
    ("T07", "B", "mixup", {"mix": "mixup", "mix_alpha": 0.2}),
    ("T08", "C", "labelsmooth", {"loss": "ls", "label_smoothing": 0.1}),
    ("T09", "C", "focal", {"loss": "focal"}),
    ("T10", "C", "ce_weighted", {"loss": "ce_weighted"}),
    ("T11", "D", "balanced", {"sampler": "balanced"}),
    ("T12", "E", "same_lr", {"lr_head": 1e-4}),
    ("T13", "E", "lr_x3", {"lr_backbone": 3e-4, "lr_head": 3e-3}),
    ("T14", "F", "ema", {"ema_decay": 0.998}),
    ("T15", "G", "res256", {"img_size": 256}),
    ("T16", "G", "epochs20", {"epochs": 20}),
]
COMBO_ID = "T20"
COMBO_AXES = ("B", "C", "D", "E", "F", "G")     # trục A (khởi tạo) không đưa vào tổ hợp
HARD = {"Chinee Apple": 0, "Snake Weed": 7}


# --------------------------------------------------------------------------- #
# Chạy và đọc kết quả
# --------------------------------------------------------------------------- #
def backbone_configs(common: dict, keys=None) -> list[Config]:
    return [Config(exp_id=eid, backbone=key, desc=key, **common) for eid, key in BACKBONES
            if keys is None or key in keys]


def t00_config(backbone: str, seed: int, common: dict, **extra) -> Config:
    return Config(exp_id="T00", backbone=backbone, desc="baseline", seed=seed, **{**common, **extra})


def training_specs(backbone: str, only=None) -> list[tuple]:
    """Các dòng của TRAINING dùng được cho backbone này (bỏ res256 với mạng cố định đầu vào)."""
    return [s for s in TRAINING if (only is None or s[0] in only)
            and not (backbone in FIXED_INPUT and "img_size" in s[3])]


def training_configs(backbone: str, common: dict, only=None) -> list[Config]:
    return [Config(exp_id=eid, backbone=backbone, desc=desc, **{**common, **over})
            for eid, _, desc, over in training_specs(backbone, only)]


def run_all(cfgs: list[Config]) -> list[dict]:
    """Chạy lần lượt; một lần hỏng (ví dụ hết bộ nhớ) được in ra và bỏ qua, các lần sau vẫn chạy."""
    out = []
    for cfg in cfgs:
        try:
            out.append(train.run(cfg))
        except Exception as e:  # noqa: BLE001 - ghi lại rồi chạy tiếp, không nuốt lỗi im lặng
            print(f"!!! [{cfg.exp_id} seed{cfg.seed}] HỎNG: {type(e).__name__}: {e}")
        if torch.cuda.is_available():
            torch.cuda.empty_cache()
    return out


def load_summary(cfg: Config) -> dict | None:
    path = train.run_dir(cfg) / "summary.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


def differs(over: dict) -> str:
    return ", ".join(f"{k}={v}" for k, v in over.items())


# --------------------------------------------------------------------------- #
# Bước 1: backbone
# --------------------------------------------------------------------------- #
def backbone_latency(keys, img_size: int = 224, device: str = "cuda", iters: int = 100) -> dict:
    """Độ trễ sơ bộ batch 1, FP32 của từng backbone (trọng số ngẫu nhiên: độ trễ không phụ thuộc trọng số)."""
    out = {}
    for key in keys:
        net = model_lib.build_model(key, pretrained=False)
        out[key] = benchmark.latency_report(net, 1, img_size, "fp32", device, iters=iters)["p50"]
        del net
    return out


def backbone_table(cfgs: list[Config], latency: dict | None = None) -> pd.DataFrame:
    rows = []
    for cfg in cfgs:
        s = load_summary(cfg)
        if s is None:
            continue
        rows.append({"exp_id": cfg.exp_id, "backbone": cfg.backbone, "weights_tag": s["weights_tag"],
                     "params_m": s["params_m"], "gmacs": s["gmacs"], "img_size": cfg.img_size,
                     "epochs": cfg.epochs, "seed": cfg.seed, "best_epoch": s["best_epoch"],
                     "val_macro_f1": s["val_macro_f1"], "val_top1": s["val_top1"],
                     "train_s_per_epoch": s["train_time_per_epoch_s"],
                     "latency_b1_ms": (latency or {}).get(cfg.backbone, np.nan), "note": "1 seed"})
    return pd.DataFrame(rows)


def pick_backbone(table: pd.DataFrame) -> str:
    """Backbone có macro-F1 val cao nhất (hòa thì lấy mạng ít GMAC hơn)."""
    best = table.sort_values(["val_macro_f1", "gmacs"], ascending=[False, True]).iloc[0]
    return str(best["backbone"])


# --------------------------------------------------------------------------- #
# Bước 2: công thức huấn luyện
# --------------------------------------------------------------------------- #
def noise_std(t00_cfgs: list[Config]) -> float:
    """std mẫu (ddof=1) của macro-F1 val qua các seed của T00: thước đo nhiễu để so sánh Δ."""
    vals = [s["val_macro_f1"] for s in map(load_summary, t00_cfgs) if s]
    return float(np.std(vals, ddof=1)) if len(vals) > 1 else float("nan")


def _verdict(delta: float, std: float) -> str:
    if not np.isfinite(std):
        return "chưa có ước lượng nhiễu"
    if delta > std:
        return "tốt hơn, vượt nhiễu (1 seed)"
    if delta < -std:
        return "kém hơn, vượt nhiễu (1 seed)"
    return "không phân biệt được với T00"


def training_table(t00_cfgs: list[Config], t_cfgs: list[Config], combo: Config | None = None,
                   combo_note: str = "") -> pd.DataFrame:
    """Bảng ablation: Δ tính với T00 CÙNG seed (seed 0); nhiễu = std của T00 qua các seed."""
    std = noise_std(t00_cfgs)
    meta = {eid: (axis, over) for eid, axis, _, over in TRAINING}
    base = next(s for c in t00_cfgs if c.seed == 0 and (s := load_summary(c)))

    def row(cfg, axis, diff, s, note):
        f1 = s["val_f1_per_class"]
        return {"exp_id": cfg.exp_id, "backbone": cfg.backbone, "axis": axis, "differs": diff, "seed": cfg.seed,
                "best_epoch": s["best_epoch"], "val_macro_f1": s["val_macro_f1"], "val_top1": s["val_top1"],
                "delta_vs_t00": s["val_macro_f1"] - base["val_macro_f1"],
                "f1_chinee_apple": f1[HARD["Chinee Apple"]], "f1_snake_weed": f1[HARD["Snake Weed"]],
                "f1_negatives": f1[8], "note": note}

    rows = []
    for cfg in t00_cfgs:
        if s := load_summary(cfg):
            rows.append(row(cfg, "-", "công thức nền", s, f"nền; std qua {len(t00_cfgs)} seed = {std:.4f}"))
    for cfg in t_cfgs:
        if s := load_summary(cfg):
            axis, over = meta[cfg.exp_id]
            rows.append(row(cfg, axis, differs(over), s, _verdict(s["val_macro_f1"] - base["val_macro_f1"], std)))
    if combo is not None and (s := load_summary(combo)):
        rows.append(row(combo, "tổ hợp", combo_note, s, _verdict(s["val_macro_f1"] - base["val_macro_f1"], std)))
    return pd.DataFrame(rows)


def combo_overrides(table: pd.DataFrame, std: float) -> tuple[dict, list[str]]:
    """Chọn tổ hợp trên VAL: mỗi trục (B-G) lấy giá trị tốt nhất nếu Δ > nhiễu.

    Nếu có ít hơn 2 trục vượt nhiễu thì lấy 2 trục có Δ lớn nhất (để vẫn kiểm tra được cộng dồn hay
    triệt tiêu, GUIDE mục 3.1). Trả về (các field ghi đè, danh sách exp_id đã dùng).
    """
    over = {eid: o for eid, _, _, o in TRAINING}
    cand = table[table["axis"].isin(COMBO_AXES)]
    best = cand.loc[cand.groupby("axis")["delta_vs_t00"].idxmax()].sort_values("delta_vs_t00", ascending=False)
    threshold = std if np.isfinite(std) else 0.0
    chosen = best[best["delta_vs_t00"] > threshold]
    if len(chosen) < 2:
        chosen = best.head(2)
    merged: dict = {}
    for eid in chosen["exp_id"]:
        merged.update(over[eid])
    return merged, chosen["exp_id"].tolist()


def pick_final(table: pd.DataFrame) -> str:
    """exp_id (seed 0) có macro-F1 val cao nhất trong T00, T01.., T20; trục A không tính."""
    cand = table[(table["seed"] == 0) & (table["axis"] != "A")]
    return str(cand.sort_values("val_macro_f1", ascending=False).iloc[0]["exp_id"])


def overrides_of(exp_id: str, combo_over: dict | None = None) -> dict:
    if exp_id == "T00":
        return {}
    if exp_id == COMBO_ID:
        return dict(combo_over or {})
    return dict(next(o for eid, _, _, o in TRAINING if eid == exp_id))


# --------------------------------------------------------------------------- #
# Bước 3: suy luận. Mọi view được tạo từ batch ảnh GỐC 256x256 (loader đánh giá ở 256).
# --------------------------------------------------------------------------- #
def _resize(size: int):
    return lambda x: x if x.shape[-1] == size else F.interpolate(
        x, size=(size, size), mode="bilinear", align_corners=False, antialias=True)


def build_views(img_size: int, fixed_input: bool) -> tuple[dict, list[dict]]:
    """Trả về (view: tên -> hàm trên batch 256, methods: list phương pháp suy luận một mô hình).

    Mỗi phương pháp: {"exp_id", "method", "views", "space", "parts"}; parts = [(img_size, số lượt)] để đo độ trễ.
    """
    s = img_size
    if s > 256:
        raise ValueError("img_size huấn luyện > 256 chưa được hỗ trợ")
    big = _resize(round(s * 256 / 224 / 4) * 4) if s == 256 else (lambda x: x)   # ảnh nền để cắt 5 crop
    center = (lambda x: x) if s == 256 else (lambda x: inference.views_multicrop(x, s)[4])
    flip = inference.view_hflip

    views = {"c": center, "c_f": lambda x: flip(center(x))}
    for i in range(5):
        views[f"crop{i}"] = lambda x, i=i: inference.views_multicrop(big(x), s)[i]
        views[f"crop{i}_f"] = lambda x, i=i: flip(inference.views_multicrop(big(x), s)[i])
    crops = [f"crop{i}" for i in range(5)]

    methods = [
        {"exp_id": "I00", "method": f"1 view (center {s})", "views": ["c"], "space": "prob", "parts": [(s, 1)]},
        {"exp_id": "I01", "method": "TTA lật ngang, gộp xác suất", "views": ["c", "c_f"], "space": "prob",
         "parts": [(s, 2)]},
        {"exp_id": "I02", "method": "TTA 5 crop, gộp xác suất", "views": crops, "space": "prob", "parts": [(s, 5)]},
        {"exp_id": "I02b", "method": "TTA 10 crop (5 crop + lật), gộp xác suất",
         "views": crops + [c + "_f" for c in crops], "space": "prob", "parts": [(s, 10)]},
        {"exp_id": "I03a", "method": "TTA lật ngang, gộp logit", "views": ["c", "c_f"], "space": "logit",
         "parts": [(s, 2)]},
        {"exp_id": "I03b", "method": "TTA 5 crop, gộp logit", "views": crops, "space": "logit", "parts": [(s, 5)]},
    ]
    if not fixed_input:
        for tag, r in zip("abc", [r for r in (256, 288, 320) if r > s]):
            views[f"r{r}"] = _resize(r)
            views[f"r{r}_f"] = lambda x, r=r: flip(_resize(r)(x))
            methods.append({"exp_id": f"I04{tag}", "method": f"1 view ở độ phân giải {r}", "views": [f"r{r}"],
                            "space": "prob", "parts": [(r, 1)]})
            methods.append({"exp_id": f"I04{tag}f", "method": f"độ phân giải {r} + lật ngang, gộp xác suất",
                            "views": [f"r{r}", f"r{r}_f"], "space": "prob", "parts": [(r, 2)]})
    return views, methods


def view_logits(net, df, images_dir, device, views: dict, names: list[str], batch_size: int = 64,
                num_workers: int = 2, amp: bool = False):
    """Chạy `net` trên mọi ảnh của `df` (đọc ở 256) với các view có tên trong `names`.

    Trả về (filenames, y_true, {tên view: logits[N, 9]}). Mặc định FP32 (amp=False).
    """
    loader = dataset.make_loader(df, images_dir, dataset.build_transforms(False, 256), batch_size, train=False,
                                 num_workers=num_workers)
    files, y, outs = inference.predict_views(net, loader, device, [views[n] for n in names], amp=amp)
    return files, y, dict(zip(names, outs))


def method_probs(logits: dict, method: dict) -> np.ndarray:
    return inference.aggregate_views([logits[v] for v in method["views"]], method["space"])


def _metric_row(y, probs) -> dict:
    m = ev.compute_metrics(y, probs.argmax(1), probs)
    return {"val_macro_f1": m["macro_f1"], "val_top1": m["top1"], "val_ece": m["ece"]}


def _softmax(logits) -> np.ndarray:
    return inference.apply_temperature(logits, 1.0)


def fit_log_temperature(probs, y) -> float:
    """Temperature scaling trên log-xác suất đã gộp: p_T = softmax(log p / T). Với 1 view, trùng TS chuẩn."""
    return inference.fit_temperature(np.log(np.clip(probs, 1e-12, None)), y)


def apply_log_temperature(probs, T: float) -> np.ndarray:
    return inference.apply_temperature(np.log(np.clip(probs, 1e-12, None)), T)


def inference_study(cfg: Config, device, members=(), ema_summary: dict | None = None, iters: int = 100,
                    max_images: int | None = None, throughput_batch: int = 32
                    ) -> tuple[pd.DataFrame, pd.DataFrame]:
    """So sánh các phương pháp suy luận trên VAL cho checkpoint tốt nhất của `cfg` (không huấn luyện lại).

    members: list (nhãn, khoá backbone, đường dẫn val_logits.npy) của các mô hình khác để ensemble.
    ema_summary: summary.json của lần chạy có EMA (T14), nếu có.
    max_images và throughput_batch khác 32 chỉ để chạy thử nhanh; giữ mặc định khi chạy thật.
    Trả về (bảng Inference, bảng Latency). Độ chính xác đo ở FP32; độ trễ theo benchmark.py.
    """
    dev = torch.device(device)
    dev_name = str(dev)
    _, val_df, _ = dataset.load_split(cfg.labels_dir, cfg.fold)
    if max_images:
        val_df = val_df.iloc[:max_images]
    net = inference.load_checkpoint(cfg.backbone, train.run_dir(cfg) / "best.pt", dev)
    model_name = f"{cfg.exp_id} seed{cfg.seed} ({cfg.backbone})"
    s = cfg.img_size
    views, methods = build_views(s, cfg.backbone in FIXED_INPUT)
    common = dict(num_workers=cfg.num_workers, batch_size=cfg.batch_size)
    _, y, logits = view_logits(net, val_df, cfg.images_dir, dev, views, list(views), **common)

    rows, lat_rows = [], []

    def latency(exp_id, parts, dtype="fp32", fused=False, label=None):
        out = {}
        for batch in (1, throughput_batch):
            r = benchmark.pipeline_latency(parts, batch, dtype, dev_name, iters=iters,
                                           config=label or exp_id, fused_bn=fused)
            lat_rows.append(r)
            out[batch] = r
        return {"p50_ms": out[1]["p50"], "p95_ms": out[1]["p95"], "p99_ms": out[1]["p99"],
                "images_per_s_b32": out[throughput_batch]["images_per_s"]}

    def add(exp_id, method, model, k, metrics, lat, single=False, note=""):
        rows.append({"exp_id": exp_id, "method": method, "model": model, "k": k, **metrics, **lat,
                     "single_model": single, "note": note})

    for m in methods:
        probs = method_probs(logits, m)
        k = sum(n for _, n in m["parts"])
        add(m["exp_id"], m["method"], model_name, k, _metric_row(y, probs),
            latency(m["exp_id"], [(net, img, n) for img, n in m["parts"]]), single=True)
    base = rows[0]
    p_base = _softmax(logits["c"])

    # I05: ensemble (trung bình xác suất 1 view của từng mô hình)
    if members and not max_images:
        m_probs = [p_base] + [_softmax(np.load(path)) for _, _, path in members]
        parts = [(net, s, 1)] + [(model_lib.build_model(key, pretrained=False), 224, 1) for _, key, _ in members]
        add("I05", "Ensemble, trung bình xác suất", " + ".join([model_name] + [lab for lab, _, _ in members]),
            len(m_probs), _metric_row(y, inference.ensemble_probs(m_probs)), latency("I05", parts),
            note="chi phí = tổng các mô hình")
        del parts

    # I06: trọng số EMA (huấn luyện ở Bước 2), không tốn thêm khi suy luận
    if ema_summary:
        add("I06", "Trọng số EMA, 1 view", f"{ema_summary['exp_id']} seed{ema_summary['seed']} (EMA)", 1,
            {"val_macro_f1": ema_summary["val_macro_f1"], "val_top1": ema_summary["val_top1"],
             "val_ece": ema_summary["val_ece"]},
            {k: base[k] for k in ("p50_ms", "p95_ms", "p99_ms", "images_per_s_b32")},
            note="mô hình khác (train có EMA); độ trễ bằng I00 vì cùng kiến trúc")

    # I07: temperature scaling, T khớp trên val
    T = fit_log_temperature(p_base, y)
    cal = _metric_row(y, apply_log_temperature(p_base, T))
    add("I07", "Temperature scaling (1 view)", model_name, 1, cal,
        {k: base[k] for k in ("p50_ms", "p95_ms", "p99_ms", "images_per_s_b32")},
        note=f"T = {T:.4f} khớp trên val; ECE val {base['val_ece']:.4f} -> {cal['val_ece']:.4f} "
             "(đo trên chính val nên lạc quan; kiểm tra thật trên test ở Bước 4)")

    # I08: gộp BN, AMP, FP16 (cùng trọng số, 1 view)
    fused = inference.fuse_conv_bn(net)
    n_bn = sum(isinstance(mod, torch.nn.BatchNorm2d) for mod in net.modules())
    if n_bn:
        err = inference.fuse_error(net, fused, s)
        _, _, lf = view_logits(fused, val_df, cfg.images_dir, dev, views, ["c"], **common)
        add("I08a", "Gộp BatchNorm vào conv (FP32)", model_name, 1, _metric_row(y, _softmax(lf["c"])),
            latency("I08a", [(fused, s, 1)], fused=True), note=f"sai số logit lớn nhất {err:.2e}")
    else:
        print(f"I08a: {cfg.backbone} không có BatchNorm, gộp BN không áp dụng")
    if dev.type == "cuda":
        _, _, la = view_logits(net, val_df, cfg.images_dir, dev, views, ["c"], amp=True, **common)
        add("I08b", "AMP (autocast)", model_name, 1, _metric_row(y, _softmax(la["c"])),
            latency("I08b", [(net, s, 1)], dtype="amp"))
        half = fused if n_bn else net
        half16 = copy.deepcopy(half).half()
        _, _, lh = view_logits(half16, val_df, cfg.images_dir, dev, {"c": lambda x: views["c"](x).half()}, ["c"],
                               **common)
        add("I08c", "FP16" + (" + gộp BN" if n_bn else ""), model_name, 1, _metric_row(y, _softmax(lh["c"])),
            latency("I08c", [(half, s, 1)], dtype="fp16", fused=bool(n_bn)))

    table = pd.DataFrame(rows)
    table["rel_cost"] = table["p50_ms"] / base["p50_ms"]
    return table, pd.DataFrame(lat_rows)


def pick_method(table: pd.DataFrame) -> str:
    """Phương pháp một-mô-hình (I00-I04) có macro-F1 val cao nhất; hòa thì lấy cách rẻ hơn."""
    cand = table[table["single_model"]]
    return str(cand.sort_values(["val_macro_f1", "k"], ascending=[False, True]).iloc[0]["exp_id"])


def pick_realtime(table: pd.DataFrame, budget_ms: float = 100.0) -> pd.Series | None:
    """Cấu hình thời gian thực: trong các cách chạy 1 lượt forward của mô hình chung kết (I00, I08*),
    lấy cách có p95 batch 1 nhỏ nhất, nếu p95 <= ngân sách."""
    cand = table[table["exp_id"].isin(["I00", "I08a", "I08b", "I08c"])].sort_values("p95_ms")
    return cand.iloc[0] if len(cand) and cand.iloc[0]["p95_ms"] <= budget_ms else None


# --------------------------------------------------------------------------- #
# Bước 4: chung kết. Test chỉ được đọc ở đây.
# --------------------------------------------------------------------------- #
def final_predictions(cfg: Config, method_id: str, device, tag: str | None = None) -> dict:
    """Ghi file dự đoán chung kết cho MỘT seed (cfg đã huấn luyện xong), test đúng một lần.

    Dùng phương pháp suy luận `method_id` (đã chọn trên val), T khớp trên val của chính seed đó:
        <tag>_seed<k>_val.csv, <tag>_seed<k>_test.csv          sau temperature scaling
        <tag>_uncal_seed<k>_test.csv                           cùng phương pháp, chưa temperature scaling
        <tag>rt_seed<k>_val.csv, <tag>rt_seed<k>_test.csv      thời gian thực: 1 view, chưa temperature scaling
    Nếu file test đã tồn tại thì KHÔNG chạy lại (README mục 2.2: test một lần mỗi seed).
    """
    tag = tag or cfg.exp_id
    pred_dir, k = Path(cfg.pred_dir), cfg.seed
    test_path = pred_dir / f"{tag}_seed{k}_test.csv"
    info_path = train.run_dir(cfg) / "final_inference.json"
    if test_path.exists():
        print(f"[{tag} seed{k}] {test_path} đã có: không chạy test lần nữa")
        return json.loads(info_path.read_text(encoding="utf-8")) if info_path.exists() else {}

    dev = torch.device(device)
    _, val_df, test_df = dataset.load_split(cfg.labels_dir, cfg.fold)
    net = inference.load_checkpoint(cfg.backbone, train.run_dir(cfg) / "best.pt", dev)
    views, methods = build_views(cfg.img_size, cfg.backbone in FIXED_INPUT)
    method = next(m for m in methods if m["exp_id"] == method_id)
    names = sorted(set(method["views"]) | {"c"})
    common = dict(num_workers=cfg.num_workers, batch_size=cfg.batch_size)

    f_val, y_val, l_val = view_logits(net, val_df, cfg.images_dir, dev, views, names, **common)
    f_test, y_test, l_test = view_logits(net, test_df, cfg.images_dir, dev, views, names, **common)
    p_val, p_test = method_probs(l_val, method), method_probs(l_test, method)
    T = fit_log_temperature(p_val, y_val)                      # chỉ dùng val

    ev.save_predictions(pred_dir / f"{tag}_seed{k}_val.csv", f_val, y_val, apply_log_temperature(p_val, T))
    ev.save_predictions(pred_dir / f"{tag}_uncal_seed{k}_test.csv", f_test, y_test, p_test)
    ev.save_predictions(pred_dir / f"{tag}rt_seed{k}_val.csv", f_val, y_val, _softmax(l_val["c"]))
    ev.save_predictions(pred_dir / f"{tag}rt_seed{k}_test.csv", f_test, y_test, _softmax(l_test["c"]))
    ev.save_predictions(test_path, f_test, y_test, apply_log_temperature(p_test, T))
    info = {"tag": tag, "seed": k, "method_id": method_id, "method": method["method"], "temperature": T}
    info_path.write_text(json.dumps(info, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"[{tag} seed{k}] đã ghi dự đoán val/test ({method['method']}, T = {T:.4f})")
    return info


def _group(pred_dir, tag: str, split: str, labels_dir):
    files = sorted(Path(pred_dir).glob(f"{tag}_seed*_{split}.csv"))
    if not files:
        return None
    return ev.load_group([str(f) for f in files], str(Path(labels_dir) / f"{split}_subset0.csv"), ref_what=split)


def final_table(pred_dir, labels_dir, configs: dict[str, str]) -> pd.DataFrame:
    """Sheet Final: mỗi seed một dòng và một dòng mean ± std cho từng tag. Chỉ số tính bằng eval.py.

    configs: {tag: mô tả cấu hình}, ví dụ {"F01": "...", "T00": "...", "F01rt": "..."}.
    """
    rows = []
    for tag, desc in configs.items():
        test, val = _group(pred_dir, tag, "test", labels_dir), _group(pred_dir, tag, "val", labels_dir)
        if test is None:
            continue
        val_f1 = {p.seed: m["macro_f1"] for p, m in zip(val.preds, val.metrics)} if val else {}
        for p, m in zip(test.preds, test.metrics):
            rows.append({"exp_id": tag, "config": desc, "seed": p.seed,
                         "val_macro_f1": val_f1.get(p.seed, np.nan), "test_macro_f1": m["macro_f1"],
                         "test_top1": m["top1"], "test_ece": m["ece"], "mean_pm_std": ""})
        f1, top1, ece = (test.summary[key] for key in ("macro_f1", "top1", "ece"))
        rows.append({"exp_id": tag, "config": desc, "seed": f"mean ({len(test.preds)} seed)",
                     "val_macro_f1": float(np.mean(list(val_f1.values()))) if val_f1 else np.nan,
                     "test_macro_f1": f1[0], "test_top1": top1[0], "test_ece": ece[0],
                     "mean_pm_std": f"macro-F1 {ev.fmt(*f1)} | top-1 {ev.fmt(*top1)} | ECE {ev.fmt(*ece)}"})
    return pd.DataFrame(rows)


def per_class_table(pred_dir, labels_dir, tags: list[str]) -> pd.DataFrame:
    """Sheet PerClass: precision/recall/F1 theo lớp trên test (mean ± std qua seed) cho từng tag."""
    rows = []
    for tag in tags:
        g = _group(pred_dir, tag, "test", labels_dir)
        if g is None:
            continue
        for i, name in enumerate(dataset.CLASS_NAMES):
            row = {"exp_id": tag, "class": name, "n_test": int(g.metrics[0]["support"][i])}
            for key in ("precision", "recall", "f1"):
                mean, std = g.summary[key]
                row[key], row[f"{key}_std"] = float(mean[i]), float(std[i])
            rows.append(row)
    return pd.DataFrame(rows)


def confusion_sum(pred_dir, labels_dir, tag: str) -> np.ndarray:
    g = _group(pred_dir, tag, "test", labels_dir)
    return sum(m["confusion"] for m in g.metrics)


# --------------------------------------------------------------------------- #
# Bước 5: bảng tổng hợp, biểu đồ, results.xlsx
# --------------------------------------------------------------------------- #
def summary_table(backbones: pd.DataFrame, training: pd.DataFrame, inference_df: pd.DataFrame,
                  final: pd.DataFrame, top: int = 10) -> pd.DataFrame:
    """Sheet Summary: top cấu hình theo macro-F1 val, kèm chi phí; luôn giữ dòng mốc T00 và chung kết."""
    rows = []
    for _, r in backbones.iterrows():
        rows.append({"exp_id": r["exp_id"], "step": "backbone", "config": r["weights_tag"],
                     "val_macro_f1": r["val_macro_f1"], "val_top1": r["val_top1"],
                     "latency_b1_p50_ms": r["latency_b1_ms"], "test_macro_f1": np.nan, "note": "1 seed"})
    for _, r in training[training["seed"] == 0].iterrows():
        rows.append({"exp_id": r["exp_id"], "step": "huấn luyện", "config": f"{r['backbone']}: {r['differs']}",
                     "val_macro_f1": r["val_macro_f1"], "val_top1": r["val_top1"],
                     "latency_b1_p50_ms": np.nan, "test_macro_f1": np.nan, "note": r["note"]})
    for _, r in inference_df.iterrows():
        rows.append({"exp_id": r["exp_id"], "step": "suy luận", "config": f"{r['method']} | {r['model']}",
                     "val_macro_f1": r["val_macro_f1"], "val_top1": r["val_top1"],
                     "latency_b1_p50_ms": r["p50_ms"], "test_macro_f1": np.nan,
                     "note": f"chi phí x{r['rel_cost']:.2f} so với I00"})
    out = pd.DataFrame(rows).sort_values("val_macro_f1", ascending=False)
    keep = out.head(top)
    if "T00" not in keep["exp_id"].values:
        keep = pd.concat([keep, out[out["exp_id"] == "T00"]])
    means = final[final["seed"].astype(str).str.startswith("mean")]
    extra = [{"exp_id": r["exp_id"], "step": "CHUNG KẾT (test)", "config": r["config"],
              "val_macro_f1": r["val_macro_f1"], "val_top1": np.nan, "latency_b1_p50_ms": np.nan,
              "test_macro_f1": r["test_macro_f1"], "note": r["mean_pm_std"]} for _, r in means.iterrows()]
    return pd.concat([pd.DataFrame(extra), keep], ignore_index=True)


COLUMNS = {   # tên cột trong results.xlsx (GUIDE mục 6.1), có đơn vị
    "exp_id": "exp_id", "backbone": "backbone", "weights_tag": "tag trọng số (timm)",
    "params_m": "#tham số (M)", "gmacs": "GMAC", "img_size": "độ phân giải (px)", "epochs": "epoch",
    "seed": "seed", "best_epoch": "epoch tốt nhất", "val_macro_f1": "macro-F1 val", "val_top1": "top-1 val",
    "train_s_per_epoch": "thời gian train/epoch (s)", "latency_b1_ms": "độ trễ batch-1 p50 (ms)",
    "note": "ghi chú", "axis": "trục thay đổi (A-G)", "differs": "khác T00 ở điểm nào",
    "delta_vs_t00": "Δ macro-F1 val so với T00 (seed 0)", "f1_chinee_apple": "F1 val Chinee Apple",
    "f1_snake_weed": "F1 val Snake Weed", "f1_negatives": "F1 val Negatives", "method": "phương pháp",
    "model": "mô hình/checkpoint", "k": "K (số view hoặc số mô hình)", "val_ece": "ECE val",
    "p50_ms": "độ trễ p50 batch-1 (ms)", "p95_ms": "độ trễ p95 batch-1 (ms)",
    "p99_ms": "độ trễ p99 batch-1 (ms)", "images_per_s_b32": "thông lượng batch-32 (ảnh/s)",
    "rel_cost": "chi phí tương đối so với I00", "single_model": "một mô hình (ứng viên chung kết)",
    "config": "cấu hình", "test_macro_f1": "macro-F1 test", "test_top1": "top-1 test", "test_ece": "ECE test",
    "mean_pm_std": "mean ± std qua seed", "class": "lớp", "n_test": "số ảnh test", "precision": "precision",
    "recall": "recall", "f1": "F1", "precision_std": "precision std", "recall_std": "recall std",
    "f1_std": "F1 std", "gpu": "GPU", "dtype": "dtype", "batch": "batch", "fused_bn": "gộp BN",
    "p50": "p50 (ms)", "p95": "p95 (ms)", "p99": "p99 (ms)", "mean": "trung bình (ms)", "n": "số lần đo",
    "images_per_s": "ảnh/s", "torch": "torch", "forwards": "số lượt forward", "step": "bước",
    "latency_b1_p50_ms": "độ trễ batch-1 p50 (ms)",
}
HIGHLIGHT = {"Backbones": "val_macro_f1", "Training": "val_macro_f1", "Inference": "val_macro_f1",
             "Summary": "val_macro_f1"}


def write_xlsx(path, sheets: dict[str, pd.DataFrame]) -> Path:
    """Ghi results.xlsx: mỗi sheet một DataFrame; cố định hàng tiêu đề, 4 chữ số thập phân, tô dòng tốt nhất."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            best = df[HIGHLIGHT[name]].astype(float).idxmax() if name in HIGHLIGHT and len(df) else None
            best_pos = df.index.get_loc(best) if best is not None else None
            df.rename(columns=COLUMNS).to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, float):
                        cell.number_format = "0.0000"
            if best_pos is not None:
                for cell in ws[best_pos + 2]:
                    cell.fill = PatternFill("solid", start_color="FFF2CC")
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    return path


def plot_tradeoff(table: pd.DataFrame, path, title: str) -> None:
    """Scatter đánh đổi macro-F1 val và độ trễ p50 batch 1 (trục x log), mỗi điểm một exp_id."""
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(9, 5.5))
    ax.scatter(table["p50_ms"], table["val_macro_f1"])
    for _, r in table.iterrows():
        ax.annotate(r["exp_id"], (r["p50_ms"], r["val_macro_f1"]), textcoords="offset points", xytext=(5, 4))
    ax.axvline(100, color="gray", linestyle="--", label="ngân sách 100 ms")
    ax.set(xscale="log", xlabel="độ trễ p50, batch 1 (ms, thang log)", ylabel="macro-F1 val", title=title)
    ax.grid(alpha=0.3)
    ax.legend()
    fig.tight_layout()
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=150)
    plt.show()


def plot_confusion(cm: np.ndarray, path, title: str) -> None:
    """Ma trận nhầm lẫn (số ảnh; hàng = nhãn thật, cột = dự đoán), màu theo tỉ lệ trong hàng."""
    import matplotlib.pyplot as plt

    norm = cm / cm.sum(axis=1, keepdims=True)
    fig, ax = plt.subplots(figsize=(9, 8))
    im = ax.imshow(norm, cmap="Blues", vmin=0, vmax=1)
    ticks = range(len(dataset.CLASS_NAMES))
    ax.set(xticks=ticks, yticks=ticks, yticklabels=dataset.CLASS_NAMES, xlabel="dự đoán", ylabel="nhãn thật",
           title=title)
    ax.set_xticklabels(dataset.CLASS_NAMES, rotation=45, ha="right")
    for i in ticks:
        for j in ticks:
            ax.text(j, i, int(cm[i, j]), ha="center", va="center", color="white" if norm[i, j] > 0.5 else "black")
    fig.colorbar(im, ax=ax, label="tỉ lệ trong hàng (recall trên đường chéo)")
    fig.tight_layout()
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=150)
    plt.show()


def config_description(cfg: Config, method: str) -> str:
    """Mô tả một dòng của cấu hình: backbone + các field khác mặc định + phương pháp suy luận."""
    default = Config()
    skip = {"exp_id", "desc", "seed", "backbone", "images_dir", "labels_dir", "out_dir", "pred_dir", "curve_dir",
            "overwrite", "save_test_predictions", "num_workers"}
    diff = {f.name: getattr(cfg, f.name) for f in dataclasses.fields(Config)
            if f.name not in skip and getattr(cfg, f.name) != getattr(default, f.name)}
    recipe = differs(diff) if diff else "công thức nền"
    return f"{cfg.backbone} | {recipe} | {method}"
